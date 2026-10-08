"""AI Resume ATS Score Checker - Streamlit + Google Gemini Flash."""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
# The first model that works is used. Override with the GEMINI_MODEL secret/env var.
DEFAULT_MODELS = ["gemini-2.5-flash", "gemini-3.5-flash", "gemini-2.0-flash"]
MAX_RESUME_CHARS = 20000
MAX_FILE_MB = 5

# Weights used to compute the overall ATS score (sum = 1.0)
WEIGHTS = {
    "keywords": 0.30,
    "formatting": 0.20,
    "sections": 0.15,
    "impact": 0.20,
    "readability": 0.15,
}

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and professional resume coach.
Evaluate the resume text you are given. If a job description is provided, judge keyword match against it;
otherwise judge against general best practice for the candidate's apparent target role.
Be honest, specific and strict. Never invent facts that are not in the resume.

Return ONLY a JSON object (no markdown, no commentary) with exactly this shape:
{
  "target_role": "string - role the resume appears to target",
  "scores": {
    "keywords": 0-100,
    "formatting": 0-100,
    "sections": 0-100,
    "impact": 0-100,
    "readability": 0-100
  },
  "summary": "2-3 sentence overall assessment",
  "strengths": ["string", ...],
  "weaknesses": ["string", ...],
  "missing_keywords": ["string", ...],
  "sections_found": ["string", ...],
  "sections_missing": ["string", ...],
  "improvements": [
    {"priority": "High|Medium|Low", "area": "string", "issue": "string", "suggestion": "string"}
  ],
  "rewrite_examples": [
    {"original": "string copied from the resume", "improved": "string"}
  ]
}

Scoring guide:
- keywords: relevant hard/soft skills, tools, job-title terms (vs. job description if given)
- formatting: ATS-friendly structure inferred from the text (clear headings, consistent dates,
  no signs of tables/columns/graphics jumbling the text)
- sections: presence of Contact, Summary, Experience, Education, Skills, etc.
- impact: quantified achievements, strong action verbs, results over duties
- readability: concise bullets, consistent tense, no typos, sensible length
Give 4-8 improvements ordered by priority and 2-4 rewrite examples that use only facts present in the resume
(use placeholders like [X%] if a number is needed)."""


# ----------------------------------------------------------------------------
# File parsing
# ----------------------------------------------------------------------------
def extract_text(uploaded_file) -> str:
    """Extract plain text from an uploaded PDF, DOCX or TXT file."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected. Please upload an unlocked copy.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = "\n".join(pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        # Resumes often use tables for layout - include them
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text


# ----------------------------------------------------------------------------
# Gemini
# ----------------------------------------------------------------------------
def get_secret(name: str):
    """Read from Streamlit secrets first, then environment variables."""
    try:
        if name in st.secrets:
            return st.secrets[name]
    except Exception:
        pass  # no secrets.toml available
    return os.environ.get(name)


def parse_json(raw: str) -> dict:
    """Parse JSON from a model response, tolerating markdown fences/extra text."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end > start:
            return json.loads(raw[start : end + 1])
        raise


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize(result: dict) -> dict:
    """Make sure every field exists and compute the overall score ourselves."""
    scores_in = result.get("scores") or {}
    scores = {k: _clamp(scores_in.get(k)) for k in WEIGHTS}
    overall = round(sum(scores[k] * w for k, w in WEIGHTS.items()))

    def as_list(key):
        value = result.get(key)
        return [str(v) for v in value] if isinstance(value, list) else []

    improvements = [i for i in (result.get("improvements") or []) if isinstance(i, dict)]
    order = {"high": 0, "medium": 1, "low": 2}
    improvements.sort(key=lambda i: order.get(str(i.get("priority", "")).lower(), 3))

    rewrites = [r for r in (result.get("rewrite_examples") or []) if isinstance(r, dict)]

    return {
        "overall": overall,
        "scores": scores,
        "target_role": str(result.get("target_role") or "Not detected"),
        "summary": str(result.get("summary") or ""),
        "strengths": as_list("strengths"),
        "weaknesses": as_list("weaknesses"),
        "missing_keywords": as_list("missing_keywords"),
        "sections_found": as_list("sections_found"),
        "sections_missing": as_list("sections_missing"),
        "improvements": improvements,
        "rewrite_examples": rewrites,
    }


def analyze_resume(resume_text: str, job_description: str, api_key: str, model_override=None) -> dict:
    client = genai.Client(api_key=api_key)
    job_description = job_description or ""

    prompt = f"RESUME TEXT:\n\"\"\"\n{resume_text[:MAX_RESUME_CHARS]}\n\"\"\"\n\n"
    if job_description.strip():
        prompt += f"JOB DESCRIPTION:\n\"\"\"\n{job_description.strip()[:8000]}\n\"\"\"\n"
    else:
        prompt += "JOB DESCRIPTION: (none provided - use general best practice)\n"

    config = types.GenerateContentConfig(
        system_instruction=SYSTEM_PROMPT,
        response_mime_type="application/json",
        temperature=0.2,
    )

    models = [model_override] if model_override else DEFAULT_MODELS
    last_error = None
    for model in models:
        try:
            response = client.models.generate_content(model=model, contents=prompt, config=config)
            return normalize(parse_json(response.text))
        except Exception as exc:  # try the next model
            last_error = exc
    raise RuntimeError(f"Gemini request failed: {last_error}")


# ----------------------------------------------------------------------------
# UI helpers
# ----------------------------------------------------------------------------
def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent - likely to pass ATS screening"
    if score >= 65:
        return "Good - a few fixes will help"
    if score >= 50:
        return "Fair - needs noticeable improvement"
    return "Poor - high risk of being filtered out"


def render_results(r: dict):
    st.divider()
    st.subheader("Your ATS score")
    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("Overall ATS score", f"{r['overall']}/100")
        st.caption(score_label(r["overall"]))
        st.caption(f"Detected target role: **{r['target_role']}**")
    with col2:
        st.progress(r["overall"] / 100)
        if r["summary"]:
            st.write(r["summary"])

    st.markdown("#### Score breakdown")
    labels = {
        "keywords": "Keywords",
        "formatting": "Formatting",
        "sections": "Sections",
        "impact": "Impact",
        "readability": "Readability",
    }
    cols = st.columns(len(labels))
    for col, (key, label) in zip(cols, labels.items()):
        col.metric(label, r["scores"][key])

    left, right = st.columns(2)
    with left:
        st.markdown("#### Strengths")
        for item in r["strengths"] or ["No strengths listed."]:
            st.markdown(f"- {item}")
    with right:
        st.markdown("#### Weaknesses")
        for item in r["weaknesses"] or ["No weaknesses listed."]:
            st.markdown(f"- {item}")

    if r["missing_keywords"]:
        st.markdown("#### Missing keywords")
        st.write(", ".join(f"`{k}`" for k in r["missing_keywords"]))

    if r["sections_missing"]:
        st.markdown("#### Missing sections")
        st.write(", ".join(r["sections_missing"]))

    st.markdown("#### Suggested improvements")
    icons = {"high": "🔴", "medium": "🟠", "low": "🟢"}
    for imp in r["improvements"]:
        pr = str(imp.get("priority", "")).strip()
        icon = icons.get(pr.lower(), "⚪")
        with st.expander(f"{icon} {pr or 'Tip'} - {imp.get('area', 'General')}"):
            st.markdown(f"**Issue:** {imp.get('issue', '')}")
            st.markdown(f"**Fix:** {imp.get('suggestion', '')}")
    if not r["improvements"]:
        st.info("No specific improvements were returned.")

    if r["rewrite_examples"]:
        st.markdown("#### Example rewrites")
        for ex in r["rewrite_examples"]:
            st.markdown(f"**Before:** {ex.get('original', '')}")
            st.markdown(f"**After:** {ex.get('improved', '')}")
            st.write("")

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(r, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def main():
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.write("Upload your resume to get an ATS score and concrete tips to improve it.")

    api_key = get_secret("GEMINI_API_KEY")
    model_override = get_secret("GEMINI_MODEL")

    with st.sidebar:
        st.header("Settings")
        if not api_key:
            api_key = st.text_input("Gemini API key", type="password",
                                    help="Get a free key at https://aistudio.google.com/app/apikey")
        else:
            st.success("API key loaded")
        st.caption("Your resume is sent to Google's Gemini API for analysis. Don't upload anything you aren't comfortable sharing.")

    uploaded = st.file_uploader("Upload resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])
    jd = st.text_area("Job description (optional, improves keyword matching)", height=150)

    if st.button("Analyze resume", type="primary"):
        if not uploaded:
            st.warning("Please upload a resume first.")
            return
        if not api_key:
            st.warning("Please provide a Gemini API key in the sidebar.")
            return
        if len(uploaded.getvalue()) > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
            return

        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded)
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return

        if len(text) < 100:
            st.error("Very little text could be extracted. If this is a scanned/image PDF, "
                     "ATS systems can't read it either - export a text-based PDF or DOCX instead.")
            return

        try:
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(text, jd, api_key, model_override)
        except Exception as exc:
            st.error(str(exc))
            return

        st.session_state["result"] = result

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
